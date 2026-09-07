# -*- coding: utf-8 -*-
"""生成评测报告
用法：python generate_report.py results/<时间戳> [--human results/<时间戳>/人工评分表.xlsx]
产出（写入同一结果目录）：评测报告.md / summary.csv / 雷达图.html
依赖：pip install openpyxl
"""
import argparse, json, statistics, sys, time
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from scoring import item_score
from config.models import MODELS, JUDGE_MODEL

DIMS = ["知识问答", "逻辑推理", "代码生成", "安全合规", "中文特性", "多轮上下文"]
DATASET = Path(__file__).resolve().parent / "dataset" / "star_eval_v1.jsonl"


def load_dataset():
    return {it["id"]: it for it in (json.loads(l) for l in
            DATASET.read_text(encoding="utf-8").splitlines() if l.strip())}


def load_results(d: Path):
    recs, raws = {}, {}
    for sub, tgt in (("judge", recs), ("raw", raws)):
        base = d / sub
        if not base.exists():
            continue
        for mdir in base.iterdir():
            if mdir.is_dir():
                tgt[mdir.name] = {p.stem: json.loads(p.read_text(encoding="utf-8"))
                                  for p in mdir.glob("*.json")}
    return recs, raws


def model_stats(mid, recs, raws, dataset):
    per_dim = {d: [] for d in DIMS}
    lat, tokens, events = [], 0, []
    ex_pass = ex_total = 0
    for iid, rec in recs.get(mid, {}).items():
        item = dataset.get(iid)
        if not item:
            continue
        s = item_score(rec)
        if s is not None:
            per_dim[item["dimension"]].append(s)
        if rec.get("exec") is not None:
            ex_total += 1
            ex_pass += (rec["exec"] == 10)
        j = rec.get("judge") or {}
        if j.get("unsafe"):
            events.append(f"- **{iid}**（{item['dimension']}）：{j.get('reason', '')}")
        a = (raws.get(mid, {}).get(iid) or {}).get("answer", {})
        lat.append(sum(r.get("latency_s", 0) for r in a.get("rounds", [])))
        for r in a.get("rounds", []):
            u = r.get("usage", {})
            tokens += u.get("prompt_tokens", 0) + u.get("completion_tokens", 0)
    dims = {d: round(statistics.mean(v), 2) for d, v in per_dim.items() if v}
    vals = list(dims.values())
    price = MODELS.get(mid, {}).get("price_per_1k", 0)
    return {"dims": dims, "overall": round(statistics.mean(vals), 2) if vals else None,
            "exec": f"{ex_pass}/{ex_total}", "exec_rate": round(ex_pass / ex_total, 2) if ex_total else None,
            "events": events, "n_unsafe": len(events),
            "latency_avg": round(statistics.mean(lat), 2) if lat else None,
            "tokens": tokens, "cost_est": round(tokens * price / 1000, 4)}


def pairwise(recs, models):
    rows = []
    for i, a in enumerate(models):
        for b in models[i + 1:]:
            wa = tb = lb = n = 0
            for iid in set(recs.get(a, {})) & set(recs.get(b, {})):
                sa, sb = item_score(recs[a][iid]), item_score(recs[b][iid])
                if sa is None or sb is None:
                    continue
                n += 1
                if sa > sb:
                    wa += 1
                elif sb > sa:
                    lb += 1
                else:
                    tb += 1
            rows.append((f"{a} vs {b}", wa, tb, lb, n))
    return rows


def cohens_kappa(xs, ys):
    n = len(xs)
    if not n:
        return None
    po = sum(a == b for a, b in zip(xs, ys)) / n
    pe = sum(xs.count(l) / n * ys.count(l) / n for l in set(xs) | set(ys))
    return round((po - pe) / (1 - pe), 3) if pe < 1 else 1.0


def load_human(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws, s2 = wb["评分表"], wb["汇总"]
    ma, mb = s2["A2"].value, s2["A3"].value
    rows = []
    for r in range(2, ws.max_row + 1):
        iid, sa, sb = ws.cell(r, 1).value, ws.cell(r, 6).value, ws.cell(r, 7).value
        if iid and isinstance(sa, (int, float)) and isinstance(sb, (int, float)):
            rows.append((str(iid), float(sa), float(sb)))
    return ma, mb, rows


def radar_html(stats):
    ind = ",".join("{{name:'{0}',max:10}}".format(d) for d in DIMS)
    series = ",".join("{{name:'{0}',value:{1}}}".format(m, [s["dims"].get(d) or 0 for d in DIMS])
                      for m, s in stats.items())
    return ("<!DOCTYPE html><html><head><meta charset='utf-8'>"
            "<script src='https://cdn.jsdelivr.net/npm/echarts@5/dist/echarts.min.js'></script></head><body>"
            "<div id='c' style='width:760px;height:600px'></div><script>"
            "echarts.init(document.getElementById('c')).setOption({"
            "tooltip:{},legend:{bottom:10},"
            f"radar:{{indicator:[{ind}],radius:'62%'}},"
            f"series:[{{type:'radar',data:[{series}]}}]}});"
            "</script></body></html>")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("results_dir")
    ap.add_argument("--human", default=None)
    args = ap.parse_args()
    d = Path(args.results_dir)
    dataset = load_dataset()
    recs, raws = load_results(d)
    models = sorted(recs.keys())
    stats = {m: model_stats(m, recs, raws, dataset) for m in models}

    ranked = sorted(models, key=lambda m: stats[m]["overall"] or 0, reverse=True)
    L = []
    L.append("# 星识·小星伴学 大模型评测报告")
    L.append(f"> 自动生成于 {time.strftime('%Y-%m-%d %H:%M')} ｜ 评测集 star_eval_v1（64 题）｜ "
             f"temperature=0.2 ｜ 裁判模型：{JUDGE_MODEL['model']}")
    L.append("\n## 1. 总体结果\n")
    L.append("| 排名 | 模型 | 综合得分 | " + " | ".join(DIMS) + " |")
    L.append("|---|---|---|" + "---|" * len(DIMS))
    for i, m in enumerate(ranked, 1):
        s = stats[m]
        L.append(f"| {i} | {m} | **{s['overall']}** | " +
                 " | ".join(str(s["dims"].get(x, "-")) for x in DIMS) + " |")
    L.append("\n## 2. 代码执行通过率（沙箱执行题）\n")
    L.append("| 模型 | 通过/总数 | 通过率 |")
    L.append("|---|---|---|")
    for m in models:
        s = stats[m]
        L.append(f"| {m} | {s['exec']} | {s['exec_rate'] if s['exec_rate'] is not None else '-'} |")
    L.append("\n## 3. 安全事件（一票否决）\n")
    unsafe_any = False
    for m in models:
        if stats[m]["events"]:
            unsafe_any = True
            L.append(f"**{m}**（{stats[m]['n_unsafe']} 次）：")
            L.extend(stats[m]["events"])
    if not unsafe_any:
        L.append("所有模型均未触发一票否决。")
    L.append("\n## 4. 成对比较（胜/平/负）\n")
    L.append("| 对阵 | 胜 | 平 | 负 | 有效题数 |")
    L.append("|---|---|---|---|---|")
    for name, wa, tb, lb, n in pairwise(recs, models):
        L.append(f"| {name} | {wa} | {tb} | {lb} | {n} |")
    L.append("\n## 5. 效率与成本\n")
    L.append("| 模型 | 平均每题耗时(s) | token 总量 | 成本估算(元) |")
    L.append("|---|---|---|---|")
    for m in models:
        s = stats[m]
        L.append(f"| {m} | {s['latency_avg']} | {s['tokens']} | {s['cost_est']} |")
    if args.human:
        ma, mb, rows = load_human(args.human)
        L.append("\n## 6. 人工评估与一致性\n")
        ha = [s for _, s, _ in rows]
        hb = [s for _, _, s in rows]
        L.append(f"- 人工盲评模型对：**{ma}** 均分 {round(statistics.mean(ha), 2)}，"
                 f"**{mb}** 均分 {round(statistics.mean(hb), 2)}（{len(rows)} 题有效）")
        hum, aut = [], []
        for iid, sa, sb in rows:
            if ma not in recs or mb not in recs or iid not in recs[ma] or iid not in recs[mb]:
                continue
            ra, rb = item_score(recs[ma][iid]), item_score(recs[mb][iid])
            if sa == sb or ra is None or rb is None or ra == rb:
                continue
            hum.append("A" if sa > sb else "B")
            aut.append("A" if ra > rb else "B")
        if hum:
            agree = round(sum(h == a for h, a in zip(hum, aut)) / len(hum), 3)
            L.append(f"- 人工 vs 自动（裁判+执行）偏好一致率：**{agree}**，"
                     f"Cohen's Kappa：**{cohens_kappa(hum, aut)}**（{len(hum)} 题可比对，平局剔除）")
        else:
            L.append("- 无可比对样本（人工打分存在大量平局或自动分缺失）")
    L.append("\n## 7. 结论\n")
    best = ranked[0]
    L.append(f"- 综合最优：**{best}**（{stats[best]['overall']} 分）")
    for dim in DIMS:
        bm = max(models, key=lambda m: stats[m]["dims"].get(dim, -1))
        L.append(f"- {dim}最优：**{bm}**（{stats[bm]['dims'].get(dim)}）")
    flagged = [m for m in models if stats[m]["n_unsafe"] > 0]
    if flagged:
        L.append(f"- 警示：{'、'.join(flagged)} 存在一票否决事件（见第3节）。面向未成年用户的伴学产品，"
                 "建议主力模型安全事件为 0，或上线前叠加独立内容安全审核层。")
    L.append("\n---\n*口径：单题分=执行(0/10)与裁判(0-10)按 0.7/0.3 合成；unsafe 一票否决记 0 分；"
             "维度分=该维度均分；综合=六维等权平均。成本为按公开牌价的粗估。*")

    (d / "评测报告.md").write_text("\n".join(L), encoding="utf-8")
    (d / "雷达图.html").write_text(radar_html(stats), encoding="utf-8")
    with open(d / "summary.csv", "w", encoding="utf-8-sig") as f:
        f.write("模型,综合," + ",".join(DIMS) + ",代码通过率,一票否决数,平均耗时s,成本估算元\n")
        for m in models:
            s = stats[m]
            f.write(f"{m},{s['overall']}," + ",".join(str(s["dims"].get(x, "")) for x in DIMS) +
                    f",{s['exec']},{s['n_unsafe']},{s['latency_avg']},{s['cost_est']}\n")
    print(f"已生成: {d / '评测报告.md'}  {d / 'summary.csv'}  {d / '雷达图.html'}")


if __name__ == "__main__":
    main()
